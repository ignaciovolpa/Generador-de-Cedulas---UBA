import streamlit as st
from docx import Document
import io
import re

st.title("Generador de Cédulas Judiciales")
st.subheader("Versión para Práctico Profesional - UBA Derecho")

# 1. Elegir tipo de cédula
tipo_cedula = st.selectbox("Tipo de Cédula", ["CIV (Capital)", "LY (Ley)"])

# 2. Campos compartidos
campos = {
    "juzgado": st.text_input("Juzgado"),
    "callejuzgado": st.text_input("Direccion del Juzgado"),
    "secretaria": st.text_input("Secretaría"),
    "fuero": st.text_input("Fuero"),
    "expediente": st.text_input("Expediente N°"),
    "caratula": st.text_input("Carátula"),
    "demandado": st.text_input("Persona a notificar"),
    "calledemandado": st.text_area("Domicilio"),
    "resolucion": st.text_area("Resolución dictada"),
    "observaciones": st.text_area("Observaciones"),
    "firmado": st.text_input("Firmado por"),
}

# 3. Campos exclusivos de la cédula LY
if tipo_cedula.startswith("LY"):
    campos.update({
        "nro_orden": st.text_input("Nro. de orden"),
        "zona": st.text_input("Zona"),
        "copias": st.text_input("Cantidad de copias"),
        "personal": st.text_input("Personal"),
        "observac": st.text_input("Observación adicional"),
        "caracter": st.text_input("Carácter (urgente, habilitación, etc.)"),
        "observaciones_especiales": st.text_area("Observaciones especiales"),
        "tipo_domicilio": st.text_input("Tipo de domicilio"),
        "sello_fuero": st.text_input("Sello del fuero"),
        "fecha": st.text_input("Fecha de emisión (por ej. 25 de abril)"),
        "anio": st.text_input("Año", value="2025")
    })

# Reemplazo de campos en el documento
def reemplazar_campos_llaves(doc, datos):
    patron = r"\{(.*?)\}"

    for p in doc.paragraphs:
        for match in re.findall(patron, p.text):
            key = match.strip().lower()
            if key in datos:
                p.text = p.text.replace("{" + match + "}", datos[key])

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for match in re.findall(patron, cell.text):
                    key = match.strip().lower()
                    if key in datos:
                        cell.text = cell.text.replace("{" + match + "}", datos[key])

# Botón
if st.button("Generar Cédula"):

    # 4. Seleccionar plantilla
    if tipo_cedula.startswith("CIV"):
        plantilla = "cedula capital_parametrizada.docx"
    else:
        plantilla = "cedula_ley_parametrizada.docx"

    # Abrir plantilla y reemplazar
    doc = Document(plantilla)
    reemplazar_campos_llaves(doc, campos)

    # Descargar
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    st.success("✅ Cédula generada correctamente.")
    st.download_button(
        label="📄 Descargar Cédula en Word",
        data=buffer,
        file_name="cedula_generada.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
